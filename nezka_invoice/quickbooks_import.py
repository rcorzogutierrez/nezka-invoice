"""Read the customer list exported from QuickBooks and turn it into clients.

QuickBooks (Desktop and Online) can export customers as an Excel file (.xlsx / .xlsm),
a CSV file or, in Desktop, an IIF file (File > Utilities > Export > Lists to IIF).
The column names differ between versions, so headers are matched against a list
of known aliases instead of fixed positions.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from nezka_invoice.models import EMAIL_RE, Client

SUPPORTED_SUFFIXES = (".csv", ".xlsx", ".xlsm", ".iif", ".txt")
EXCEL_SUFFIXES = (".xlsx", ".xlsm")

# Normalized header -> client field. Headers are lowercased and stripped of
# everything that is not a letter or digit ("Bill to 1" -> "billto1").
ALIASES = {
    "name": ("customer", "customername", "name", "displayname", "customerfullname", "customerdisplayname", "fullname"),
    "company": ("company", "companyname"),
    "contact_name": ("contact", "primarycontact", "cont1", "contactname"),
    "first_name": ("firstname",),
    "last_name": ("lastname",),
    "email": ("email", "mainemail", "emailaddress", "primaryemail"),
    "phone": ("phone", "mainphone", "phone1", "phonenumber", "workphone", "primaryphone"),
    "mobile": ("mobile", "mobilephone", "cell", "cellphone"),
    "alt_phone": ("altphone", "phone2", "alternatephone", "otherphone"),
    "street": ("streetaddress", "billingstreet", "billingaddress", "billtostreet1", "billingaddressline1", "address"),
    "street2": ("billtostreet2", "billingaddressline2"),
    "city": ("city", "billtocity", "billingcity", "billingaddresscity"),
    "state": ("state", "billtostate", "billingstate", "billingaddressstate", "province", "stateprovince"),
    "zip_code": ("zip", "zipcode", "postalcode", "billtozip", "billingzip", "billingzipcode", "billingaddresspostalcode"),
    "country": ("country", "billtocountry", "billingcountry", "billingaddresscountry"),
    "tax_id": ("taxid", "resalenum", "resalenumber", "taxregistrationnumber", "taxresaleno", "ein"),
    "fax": ("fax", "faxnum"),
    "secondary_contact": ("secondarycontact", "altcontact", "cont2"),
    "terms": ("terms",),
    "account_no": ("accountno", "accountnumber"),
    "notes": ("note", "notes", "notepad"),
    "status": ("activestatus", "status", "active"),
    "hidden": ("hidden",),
}
# Free-form address lines: "Bill to 1".."Bill to 5" (Desktop) and BADDR1..5 (IIF).
ADDRESS_LINE_RE = re.compile(r"^(?:billto|baddr|billingaddress)([1-5])$")
SHIP_LINE_RE = re.compile(r"^(?:shipto|saddr)([1-5])$")

HEADER_FIELDS = {alias: field_name for field_name, aliases in ALIASES.items() for alias in aliases}

# The "Bill to" lines are typed by hand, so the last line comes in many shapes:
# "Miami, FL 33142", "MIAMI., FL. 33142", "Coral Gables,FL 33146", "HIALEAH,FL.33016-6314", "MIAMI FLO,33142".
STATE = r"(?P<state>FLORIDA|FLA|FLO|[A-Z]{2})"
CITY_STATE_ZIP_RE = re.compile(
    rf"^(?P<city>.*?)[\s,.]*\b{STATE}(?![a-z])[\s,.]*(?P<zip>\d{{5}}(?:-\d{{4}})?)(?!\d)\W?[a-z0-9]*$", re.IGNORECASE
)
CITY_STATE_RE = re.compile(rf"^(?P<city>[^\d,]+?)\s*,\s*{STATE}\.?$", re.IGNORECASE)
FLORIDA = {"FLORIDA", "FLA", "FLO"}
COUNTRIES = {"usa", "us", "u.s.a", "united states", "united states of america", "estados unidos"}
ZIP_ONLY_RE = re.compile(r"^[\s,.]*(?:(?:FLORIDA|FLA|FLO|FL)\.?)?[\s,.]*\d{5}(?:-\d{4})?\W*$", re.IGNORECASE)
# Cities that are typed on the same line as the street ("2101 NW 15 AVE MIAMI, FL 33142").
KNOWN_CITIES = (
    "north miami beach", "north miami", "miami beach", "miami gardens", "miami lakes", "miami",
    "hialeah gardens", "hialeah", "doral", "medley", "miramar", "hollywood", "homestead", "davie",
    "opa locka", "opa-locka", "sunrise", "pompano beach", "fort lauderdale", "ft lauderdale", "ft. lauderdale",
    "fort pierce", "fort. pierce", "boca raton", "coral gables", "kendall", "sweetwater",
)
SAME_AS_BILLING = {"same", "same as above", "igual"}


def normalize_header(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


@dataclass
class ImportRow:
    """One customer from the file, ready to be saved, or the reason it is skipped."""

    client: Client
    line: int
    skip_reason: str = ""

    @property
    def importable(self) -> bool:
        return not self.skip_reason


@dataclass
class ImportPreview:
    rows: list[ImportRow] = field(default_factory=list)
    jobs_skipped: int = 0

    @property
    def new_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if r.importable]

    @property
    def skipped_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if not r.importable]


class ImportFileError(ValueError):
    pass


# Reading files --------------------------------------------------------------


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_csv(data: bytes) -> list[list[str]]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl is in requirements.txt
        raise ImportFileError("Falta el módulo openpyxl para leer archivos de Excel.") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(f"No se pudo abrir el archivo de Excel: {exc}") from exc
    try:
        # QuickBooks' workbook can include an "Export Tips" sheet; use the first sheet with a customer header.
        sheets = [workbook.active] + [s for s in workbook.worksheets if s is not workbook.active]
        tables = []
        for sheet in sheets:
            rows = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            try:
                _find_header(rows)
            except ImportFileError:
                tables.append(rows)
                continue
            return rows
        return tables[0] if tables else []
    finally:
        workbook.close()


def _read_iif(data: bytes) -> list[list[str]]:
    """Keep only the customer section (!CUST header + CUST rows) of an IIF file."""
    rows = []
    for line in _decode(data).splitlines():
        cells = line.split("\t")
        tag = cells[0].strip().upper()
        if tag == "!CUST":
            rows.append(cells[1:])
        elif tag == "CUST" and rows:
            rows.append([c.strip('"') for c in cells[1:]])
    if not rows:
        raise ImportFileError("El archivo IIF no tiene una lista de clientes (sección !CUST).")
    return rows


def read_rows(path: str | Path) -> list[list[str]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return _read_xlsx(path)
    if suffix == ".xls":
        raise ImportFileError(
            "Los archivos .xls (Excel 97-2003) no se pueden leer. "
            "En QuickBooks exporta como CSV o guárdalo en Excel como .xlsx."
        )
    data = path.read_bytes()
    if suffix == ".iif" or data.lstrip().upper().startswith(b"!CUST") or b"\n!CUST" in data.upper():
        return _read_iif(data)
    return _read_csv(data)


# Mapping --------------------------------------------------------------------


def _find_header(rows: list[list[str]]) -> tuple[int, dict[int, str]]:
    """Return the index of the header row and a map column -> field name.

    QuickBooks reports put a title and blank lines above the header, so the
    first row that has a customer-name column is used.
    """
    for index, row in enumerate(rows[:30]):
        mapping = {}
        for column, header in enumerate(row):
            key = normalize_header(header)
            if key in HEADER_FIELDS:
                mapping.setdefault(column, HEADER_FIELDS[key])
            elif match := ADDRESS_LINE_RE.match(key):
                mapping[column] = f"line{match.group(1)}"
            elif match := SHIP_LINE_RE.match(key):
                mapping[column] = f"ship{match.group(1)}"
        if "name" in mapping.values() or "company" in mapping.values():
            return index, mapping
    raise ImportFileError(
        "No se encontró la columna con el nombre del cliente (Customer / Name). "
        "Revisa que sea la lista de clientes exportada desde QuickBooks."
    )


def _clean(value) -> str:
    text = "" if value is None else re.sub(r"[ \t]{2,}", " ", str(value).strip())
    return "" if text.lower() in ("nan", "none") else text


def _split_address_lines(lines: list[str], skip: set[str]) -> dict[str, str]:
    """Split QuickBooks' free-form "Bill to" lines into address / city / state / zip / country."""
    # The first line is usually the customer's name and the second the contact; drop those.
    lines = [line for line in lines if line and normalize_header(line) not in skip]
    # A ZIP typed on its own line belongs to the line above ("MIAMI FL," / "33142").
    merged: list[str] = []
    for line in lines:
        if merged and ZIP_ONLY_RE.match(line) and not CITY_STATE_ZIP_RE.match(merged[-1]):
            merged[-1] = f"{merged[-1]} {line.strip()}"
        else:
            merged.append(line)
    lines = merged
    result = {}
    for index in range(len(lines) - 1, -1, -1):
        match = CITY_STATE_ZIP_RE.match(lines[index]) or CITY_STATE_RE.match(lines[index])
        if not match:
            continue
        city = match["city"].strip(" ,.")
        if not city and index > 0 and CITY_STATE_ZIP_RE.match(lines[index - 1]):
            continue  # "MIAMI, FL33178" followed by a repeated "FL 33178"
        if any(ch.isdigit() for ch in city):
            # The street is on the same line: "2101 NW 15 AVE MIAMI, FL 33142" or "1335 NW 21 TERRACE FL 33142".
            street = city
            city = ""
            for known in KNOWN_CITIES:
                if street.lower().endswith(" " + known):
                    street, city = street[: -len(known)].rstrip(" ,."), street[-len(known) :]
                    break
            result["street_tail"] = street
        state = match["state"].upper()
        result["city"] = city.title() if city.isupper() else city
        result["state"] = "FL" if state in FLORIDA else state
        if "zip" in match.groupdict():
            result["zip_code"] = match["zip"]
        extra = []
        for line in lines[index + 1 :]:
            cleaned = line.strip(" .`'").lower()
            if cleaned.replace(".", "") in {c.replace(".", "") for c in COUNTRIES}:
                result["country"] = "USA"
            elif cleaned not in ("florida", "fl") and not CITY_STATE_ZIP_RE.match(line):
                extra.append(line)
        lines = lines[:index] + ([result.pop("street_tail")] if "street_tail" in result else []) + extra
        break
    if lines:
        result["address"] = ", ".join(lines)
    return result


def _first_email(text: str) -> tuple[str, str]:
    """Return (first valid email, leftover text). QuickBooks allows several separated by , or ;."""
    parts = [p.strip() for p in re.split(r"[;,\s]+", text) if p.strip()]
    valid = [p for p in parts if EMAIL_RE.match(p)]
    if not valid:
        return "", text
    others = [p for p in parts if p != valid[0]]
    return valid[0], ", ".join(others)


def row_to_client(values: dict[str, str]) -> Client:
    name = values.get("name", "")
    company = values.get("company", "")
    if not name:
        name = company
    contact = values.get("contact_name", "") or " ".join(
        p for p in (values.get("first_name", ""), values.get("last_name", "")) if p
    )
    notes = [values.get("notes", "")]

    email, extra_email = _first_email(values.get("email", ""))
    if extra_email:
        if email:
            notes.append(f"Otros correos: {extra_email}")
        elif "@" not in extra_email:
            notes.append(f"Sitio web: {extra_email}")
        else:
            notes.append(f"Correo no válido en QuickBooks: {extra_email}")

    phone = values.get("phone", "")
    mobile = values.get("mobile", "") or values.get("alt_phone", "")
    if values.get("mobile") and values.get("alt_phone"):
        notes.append(f"Otro teléfono: {values['alt_phone']}")
    if values.get("fax"):
        notes.append(f"Fax: {values['fax']}")
    if values.get("secondary_contact"):
        notes.append(f"Otro contacto: {values['secondary_contact']}")
    if values.get("terms"):
        notes.append(f"Términos de pago: {values['terms']}")
    if values.get("account_no"):
        notes.append(f"Cuenta: {values['account_no']}")
    ship = [values.get(f"ship{i}", "") for i in range(1, 6)]
    ship = [line for line in ship if line and normalize_header(line) not in {normalize_header(name), normalize_header(company)}]
    if ship and not all(line.strip(" .`'").lower() in SAME_AS_BILLING for line in ship):
        notes.append("Enviar a: " + ", ".join(line for line in ship if line.strip(" .`'").lower() not in SAME_AS_BILLING))

    address = {}
    lines = [values.get(f"line{i}", "") for i in range(1, 6)]
    if any(lines):
        skip = {normalize_header(t) for t in (name, company, contact)} - {""}
        address = _split_address_lines(lines, skip)
    street = ", ".join(p for p in (values.get("street", ""), values.get("street2", "")) if p)
    if street:
        address["address"] = street
    for key in ("city", "state", "zip_code", "country"):
        if values.get(key):
            address[key] = values[key]

    status = values.get("status", "").lower()
    active = not (status.startswith("inactiv") or status in ("no", "false", "0") or values.get("hidden", "").upper() == "Y")

    return Client(
        name=name,
        branch=company if company and company.lower() != name.lower() else "",
        contact_name=contact,
        email=email,
        phone=phone,
        mobile=mobile,
        address=address.get("address", ""),
        city=address.get("city", ""),
        state=address.get("state", ""),
        zip_code=address.get("zip_code", ""),
        country=address.get("country", ""),
        tax_id=values.get("tax_id", ""),
        notes="\n".join(n for n in notes if n),
        active=active,
    ).normalized()


def parse_rows(rows: list[list[str]], existing_names: set[str] = frozenset()) -> ImportPreview:
    """Turn raw rows into an ImportPreview, marking duplicates and QuickBooks jobs."""
    header_index, mapping = _find_header(rows)
    existing = {n.strip().lower() for n in existing_names}
    seen: set[str] = set()
    preview = ImportPreview()
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values: dict[str, str] = {}
        for column, field_name in mapping.items():
            if column < len(row) and (text := _clean(row[column])):
                values.setdefault(field_name, text)
        if not values:
            continue
        name = values.get("name", "")
        # "Customer:Job" rows are jobs of a customer in QuickBooks Desktop, not separate customers.
        if ":" in name:
            preview.jobs_skipped += 1
            continue
        # Report exports end with a "TOTAL" row; customers like "TOTAL GREEN TROPICALS" are kept.
        if name.lower() == "total":
            continue
        client = row_to_client(values)
        key = client.name.lower()
        if not client.name:
            reason = "Sin nombre"
        elif key in existing:
            reason = "Ya existe"
        elif key in seen:
            reason = "Repetido en el archivo"
        else:
            reason = ""
        seen.add(key)
        preview.rows.append(ImportRow(client, offset, reason))
    return preview


def preview_file(path: str | Path, existing_names: set[str] = frozenset()) -> ImportPreview:
    return parse_rows(read_rows(path), existing_names)


def import_clients(repo, preview: ImportPreview) -> list[Client]:
    """Save every importable row; returns the saved clients."""
    saved = []
    for row in preview.new_rows:
        saved.append(repo.add(row.client))
    return saved
